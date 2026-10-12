"""Import DFØ's open Doffin dataset ("Kunngjøringer av offentlige anskaffelser", data.norge.no) into the local database.

The dataset is published as files per year (CSV; spreadsheets are accepted too). Column names are matched against a
list of known Norwegian/English headers; nothing is guessed:

* a notice needs a notice id column;
* a company is linked to a notice only through an organisation-number column whose role is known (buyer =
  oppdragsgiver, winner = leverandør). A bare "organisasjonsnummer" column has no role, so it is ignored unless the
  operator maps it explicitly (``--map buyer_org="Organisasjonsnummer"``);
* names without organisation numbers are never used to link a company.

Every imported notice keeps the dataset name and its period (from ``--period``, or the span of publication dates in
the file), so profiles can say which period the facts come from.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from ..core.util import clean_ws, now_iso, sha256_text
from ..identity.employer import org_numbers_in
from ..identity.normalization import ascii_fold
from ..storage.db import Database

ALIASES: dict[str, set[str]] = {
    "notice_id": {"doffin referanse", "doffinreferanse", "doffin id", "doffin referansenummer", "kunngjoringsnummer", "kunngjoring id", "kunngjoringsid",
                  "notice id", "noticeid", "notice number", "referanse", "referansenummer", "id"},
    "title": {"tittel", "title", "navn pa anskaffelsen", "anskaffelsens navn", "anskaffelsens tittel", "kunngjoringstittel", "notice title", "navn"},
    "notice_type": {"kunngjoringstype", "type kunngjoring", "notice type", "skjematype", "skjema", "type", "kunngjoringsskjema"},
    "published": {"publiseringsdato", "publisert", "publisert dato", "dato publisert", "publication date", "published", "kunngjort", "kunngjort dato",
                  "publiseringstidspunkt"},
    "buyer_name": {"oppdragsgiver", "oppdragsgivers navn", "navn oppdragsgiver", "oppdragsgiver navn", "buyer", "buyer name", "innkjoper"},
    "buyer_org": {"oppdragsgivers organisasjonsnummer", "oppdragsgiver organisasjonsnummer", "organisasjonsnummer oppdragsgiver", "oppdragsgiver orgnr",
                  "oppdragsgivers orgnr", "orgnr oppdragsgiver", "oppdragsgiver org nr", "buyer organisation number", "buyer organization number", "buyer org number",
                  "buyer orgnr", "buyer national id", "buyer identifier"},
    "winner_name": {"leverandor", "leverandorens navn", "leverandor navn", "vinner", "vinnende leverandor", "tildelt leverandor", "winner", "winner name",
                    "supplier", "supplier name", "kontraktspart"},
    "winner_org": {"leverandorens organisasjonsnummer", "leverandor organisasjonsnummer", "organisasjonsnummer leverandor", "leverandor orgnr", "leverandorens orgnr",
                   "vinner organisasjonsnummer", "vinnerens organisasjonsnummer", "tildelt leverandor orgnr", "winner organisation number",
                   "winner organization number", "winner org number", "winner identifier", "supplier organisation number", "supplier organization number",
                   "supplier orgnr"},
}
ROLE_FIELDS = {"buyer": ("buyer_org", "buyer_name"), "winner": ("winner_org", "winner_name")}
BARE_ORG_HEADERS = {"organisasjonsnummer", "orgnr", "org nr", "organisation number", "organization number"}


def header_key(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", ascii_fold(header or "")).strip()


@dataclass
class ParsedNotice:
    notice_id: str
    title: str | None
    notice_type: str | None
    published: str | None
    buyer_name: str | None
    parties: list[tuple[str, str, str | None]]  # (org_number, role, name)
    row_sha256: str

    @property
    def url(self) -> str | None:
        return f"https://doffin.no/notices/{self.notice_id}" if re.fullmatch(r"\d{4}-\d{3,}", self.notice_id) else None


@dataclass
class ParseResult:
    columns: dict[str, str] = field(default_factory=dict)  # field → header used
    unrecognized: list[str] = field(default_factory=list)
    notices: list[ParsedNotice] = field(default_factory=list)
    rows: int = 0
    skipped: dict[str, int] = field(default_factory=dict)
    problems: list[str] = field(default_factory=list)

    def skip(self, reason: str) -> None:
        self.skipped[reason] = self.skipped.get(reason, 0) + 1


def _date(value: str | None) -> str | None:
    text = clean_ws(value)
    if not text:
        return None
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", text)
    if m:
        return m.group(0)
    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", text)
    if m:
        try:
            return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))).date().isoformat()
        except ValueError:
            return None
    return None


def map_columns(headers: list[str], explicit: dict[str, str] | None = None) -> tuple[dict[str, str], list[str], list[str]]:
    """→ (field → header, unrecognized headers, problems)."""
    columns: dict[str, str] = {}
    problems: list[str] = []
    by_key = {header_key(h): h for h in headers}
    for fld, header in (explicit or {}).items():
        if fld not in ALIASES:
            problems.append(f"unknown field in --map: {fld}")
        elif header not in headers and header_key(header) not in by_key:
            problems.append(f"--map {fld}: no column named {header!r}")
        else:
            columns[fld] = header if header in headers else by_key[header_key(header)]
    for fld, aliases in ALIASES.items():
        if fld in columns:
            continue
        hit = next((h for h in headers if header_key(h) in aliases), None)
        if hit and hit not in columns.values():
            columns[fld] = hit
    used = set(columns.values())
    bare = [h for h in headers if header_key(h) in BARE_ORG_HEADERS and h not in used]
    if bare and not ({"buyer_org", "winner_org"} & set(columns)):
        problems.append(f"column {bare[0]!r} has no role (buyer or supplier); map it with --map buyer_org=... or --map winner_org=... — it is ignored")
    return columns, [h for h in headers if h not in used], problems


def parse_rows(headers: list[str], rows: list[list[str]], explicit: dict[str, str] | None = None) -> ParseResult:
    out = ParseResult()
    out.columns, out.unrecognized, out.problems = map_columns(headers, explicit)
    if "notice_id" not in out.columns:
        out.problems.append("no notice-id column recognised; nothing can be imported")
        return out
    if not ({"buyer_org", "winner_org"} & set(out.columns)):
        out.problems.append("no organisation-number column with a known role; notices cannot be linked to companies")
    idx = {fld: headers.index(h) for fld, h in out.columns.items()}

    def cell(row: list[str], fld: str) -> str | None:
        i = idx.get(fld)
        return clean_ws(row[i]) if i is not None and i < len(row) and clean_ws(row[i]) else None

    for row in rows:
        out.rows += 1
        notice_id = cell(row, "notice_id")
        if not notice_id:
            out.skip("no notice id")
            continue
        parties: list[tuple[str, str, str | None]] = []
        for role, (org_f, name_f) in ROLE_FIELDS.items():
            for org in sorted(org_numbers_in([cell(row, org_f)])):
                parties.append((org, role, cell(row, name_f)))
        if not parties:
            out.skip("no organisation number for a buyer or supplier")
            continue
        record = {h: (row[i] if i < len(row) else "") for i, h in enumerate(headers)}
        out.notices.append(ParsedNotice(notice_id=notice_id, title=cell(row, "title"), notice_type=cell(row, "notice_type"), published=_date(cell(row, "published")),
                                        buyer_name=cell(row, "buyer_name"), parties=parties, row_sha256=sha256_text(repr(sorted(record.items())))))
    return out


def read_table(path: Path) -> tuple[list[str], list[list[str]]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
        return (rows[0], rows[1:]) if rows else ([], [])
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        delimiter = csv.Sniffer().sniff(text[:8192], delimiters=";,\t|").delimiter
    except csv.Error:
        delimiter = ";" if text[:2000].count(";") > text[:2000].count(",") else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return (rows[0], rows[1:]) if rows else ([], [])


def dataset_period(notices: list[ParsedNotice], explicit: str | None, filename: str) -> str:
    if explicit:
        return explicit
    years = sorted({n.published[:4] for n in notices if n.published})
    if years:
        return years[0] if years[0] == years[-1] else f"{years[0]}–{years[-1]}"
    m = re.search(r"(20\d{2})", filename)
    return m.group(1) if m else "unknown period"


def store(db: Database, parsed: ParseResult, *, dataset: str, period: str) -> dict[str, int]:
    imported = now_iso()
    notices = parties = 0
    with db.transaction() as conn:
        old = [r[0] for r in conn.execute("SELECT notice_key FROM procurement_notices WHERE source = 'doffin_open_data' AND dataset = ?", (dataset,)).fetchall()]
        for key in old:
            conn.execute("DELETE FROM procurement_parties WHERE notice_key = ?", (key,))
        conn.execute("DELETE FROM procurement_notices WHERE source = 'doffin_open_data' AND dataset = ?", (dataset,))
        for n in parsed.notices:
            key = f"doffin:{n.notice_id}"
            conn.execute("INSERT OR REPLACE INTO procurement_notices(notice_key, source, notice_id, title, notice_type, published_at, buyer_name, url, dataset, "
                         "dataset_period, row_sha256, imported_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                         (key, "doffin_open_data", n.notice_id, n.title, n.notice_type, n.published, n.buyer_name, n.url, dataset, period, n.row_sha256, imported))
            notices += 1
            for org, role, name in n.parties:
                conn.execute("INSERT OR IGNORE INTO procurement_parties(notice_key, org_number, role, name) VALUES (?,?,?,?)", (key, org, role, name))
                parties += 1
    return {"notices": notices, "parties": parties}


def import_file(db: Database, path: Path, *, period: str | None = None, explicit: dict[str, str] | None = None, dry_run: bool = False) -> dict[str, Any]:
    headers, rows = read_table(path)
    parsed = parse_rows(headers, rows, explicit)
    per = dataset_period(parsed.notices, period, path.name)
    stored = {"notices": 0, "parties": 0}
    if not dry_run and parsed.notices:
        stored = store(db, parsed, dataset=path.name, period=per)
    return {"file": path.name, "period": per, "rows": parsed.rows, "linkable_notices": len(parsed.notices), "stored": stored, "columns": parsed.columns,
            "unrecognized_columns": parsed.unrecognized, "skipped": parsed.skipped, "problems": parsed.problems}
